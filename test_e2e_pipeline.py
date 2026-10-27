"""
End-to-End Verification of Master Spreadsheet Management System using realistic sample files.
"""
from pathlib import Path
import openpyxl
from core.config import DATA_DIR, MASTER_FILE_PATH, ensure_directories
from core.excel_exporter import export_dataframe_to_csv, write_master_excel_workbook
from core.file_loader import load_file_datasets
from core.master_manager import MasterManager
from core.workflow import ConsolidationWorkflow

def run_e2e():
    print("--- 1. Initializing Master Manager and Directories ---")
    ensure_directories()
    manager = MasterManager()
    if manager.exists():
        manager.master_path.unlink()

    sample_dir = Path(__file__).parent / "sample_data"
    excel_file = sample_dir / "companies_batch1.xlsx"
    csv1_file = sample_dir / "customers.csv"
    csv2_file = sample_dir / "partners_semicolon.csv"

    print("--- 2. Loading Files and Inspecting Worksheets/CSVs ---")
    ds_excel = load_file_datasets(excel_file)
    print(f"Loaded from Excel: {len(ds_excel)} sheets: {[d.sheet_name for d in ds_excel]}")
    for d in ds_excel:
        print(f"  - Sheet '{d.sheet_name}': {d.row_count} rows, cols: {d.columns}")

    ds_csv1 = load_file_datasets(csv1_file)
    print(f"Loaded from CSV 1: {ds_csv1[0].row_count} rows, cols: {ds_csv1[0].columns}")

    ds_csv2 = load_file_datasets(csv2_file)
    print(f"Loaded from CSV 2 (semicolon): {ds_csv2[0].row_count} rows, cols: {ds_csv2[0].columns}")

    all_datasets = ds_excel + ds_csv1 + ds_csv2

    print("--- 3. Initializing Workflow and Schema Alignment ---")
    workflow = ConsolidationWorkflow(manager)
    workflow.set_datasets(all_datasets)

    print("Target Canonical Columns:", workflow.schema_mapper.target_columns)
    for ds in all_datasets:
        ds_id = f"{ds.source_filename}::{ds.sheet_name}"
        print(f"Mappings for {ds_id}:", workflow.schema_mapper.mappings.get(ds_id))

    print("--- 4. Building Consolidated Staging DataFrame ---")
    staging_df = workflow.build_consolidated_staging(include_existing_master=False)
    print(f"Staging records: {len(staging_df)} rows")

    print("--- 5. Duplicate Detection (Exact & Probable) ---")
    exact_g, prob_g = workflow.detect_duplicates(
        comparison_mode="selected_columns",
        selected_columns=["Company Name"],
        enable_fuzzy=True,
        fuzzy_columns=["Company Name"],
        fuzzy_threshold=80.0,
    )
    print(f"Exact Duplicate Groups: {len(exact_g)}")
    for g in exact_g:
        print(f"  Exact Group #{g.group_id}: {len(g.row_indices)} rows, matching: {g.matching_fields}")
        g.resolution = "merge_complementary"

    print(f"Probable Duplicate Groups: {len(prob_g)}")
    for g in prob_g:
        print(f"  Probable Group #{g.group_id} ({g.confidence}%): {g.explanation}")

    print("--- 6. Applying Resolutions & Merging ---")
    cleaned_df, review_df = workflow.apply_resolutions()
    print(f"Cleaned records: {len(cleaned_df)} rows")
    print(f"Summary metrics: {workflow.processing_summary}")

    print("--- 7. Committing to Master.xlsx ---")
    ok, msg, backup = workflow.commit_to_master()
    print(f"Commit result: {ok}, message: {msg}")

    print("--- 8. Verifying Master.xlsx Structure & Worksheets ---")
    wb = openpyxl.load_workbook(MASTER_FILE_PATH)
    print("Worksheets in Master.xlsx:", wb.sheetnames)
    assert "Cleaned_Master" in wb.sheetnames
    assert "Duplicate_Report" in wb.sheetnames
    assert "Processing_Summary" in wb.sheetnames
    assert "Review_Required" in wb.sheetnames
    assert "Source_Trace" in wb.sheetnames

    print("--- 9. Verifying CSV Export ---")
    csv_out = export_dataframe_to_csv(cleaned_df)
    print(f"CSV output length: {len(csv_out)} chars, sample:\n{csv_out[:200]}...")

    print("--- 10. Verifying Repeated Upload Idempotency ---")
    workflow2 = ConsolidationWorkflow(manager)
    workflow2.set_datasets([all_datasets[0]])  # Re-upload first dataset
    workflow2.build_consolidated_staging(include_existing_master=True)
    workflow2.detect_duplicates(comparison_mode="selected_columns", selected_columns=["Company Name"])
    cleaned2, _ = workflow2.apply_resolutions()
    print(f"Rows after repeated upload deduplication: {len(cleaned2)} (should match previous cleaned rows: {len(cleaned_df)})")
    assert len(cleaned2) == len(cleaned_df)

    print("\n[SUCCESS] ALL END-TO-END WORKFLOW VERIFICATIONS PASSED SUCCESSFULLY!")

if __name__ == "__main__":
    run_e2e()

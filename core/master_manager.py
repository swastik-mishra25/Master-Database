"""
Master workbook persistence, backups, recovery, and updates.
"""
import os
import shutil
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
import pandas as pd
import openpyxl

from .config import BACKUP_DIR, DATA_DIR, MASTER_FILE_PATH, MASTER_FILENAME, TEMP_DIR, ensure_directories
from .file_loader import read_excel_dataset
from .normalizer import PROVENANCE_COLS


class MasterManager:
    """Manages the lifecycle of Master.xlsx, including backups and safe atomic updates."""

    def __init__(self, master_path: Path = MASTER_FILE_PATH, backup_dir: Path = BACKUP_DIR):
        ensure_directories()
        self.master_path = master_path
        self.backup_dir = backup_dir

    def exists(self) -> bool:
        """Check if Master.xlsx exists and is non-empty."""
        return self.master_path.exists() and self.master_path.stat().st_size > 0

    def load_master(self) -> Optional[pd.DataFrame]:
        """Load the master dataset from Master.xlsx."""
        if not self.exists():
            return None

        try:
            # Check sheet names
            wb = openpyxl.load_workbook(self.master_path, read_only=True)
            sheet_names = wb.sheetnames
            wb.close()

            # Target Cleaned_Master if available, otherwise first sheet
            target_sheet = "Cleaned_Master" if "Cleaned_Master" in sheet_names else sheet_names[0]
            dataset = read_excel_dataset(self.master_path, filename=MASTER_FILENAME, sheet_name=target_sheet)
            return dataset.df
        except Exception as e:
            raise RuntimeError(f"Failed to load Master workbook: {str(e)}")

    def get_summary(self) -> Dict[str, Any]:
        """Get summary metadata of the current Master.xlsx."""
        if not self.exists():
            return {
                "exists": False,
                "row_count": 0,
                "column_count": 0,
                "columns": [],
                "file_size_kb": 0.0,
                "last_modified": None,
                "path": str(self.master_path),
            }

        stat = self.master_path.stat()
        df = self.load_master()
        cols = [c for c in df.columns if c not in PROVENANCE_COLS] if df is not None else []
        row_count = len(df) if df is not None else 0

        return {
            "exists": True,
            "row_count": row_count,
            "column_count": len(cols),
            "columns": cols,
            "file_size_kb": round(stat.st_size / 1024, 2),
            "last_modified": datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M:%S"),
            "path": str(self.master_path),
        }

    def create_backup(self, prefix: str = "Master_backup") -> Optional[Path]:
        """Create a timestamped backup of the current Master.xlsx."""
        if not self.exists():
            return None

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        backup_filename = f"{prefix}_{timestamp}.xlsx"
        backup_path = self.backup_dir / backup_filename

        counter = 1
        while backup_path.exists():
            backup_path = self.backup_dir / f"{prefix}_{timestamp}_{counter}.xlsx"
            counter += 1

        try:
            shutil.copy2(self.master_path, backup_path)
            return backup_path
        except Exception as e:
            raise RuntimeError(f"Failed to create backup: {str(e)}")

    def save_master(
        self,
        df: pd.DataFrame,
        create_backup: bool = True,
        extra_sheets: Optional[Dict[str, pd.DataFrame]] = None,
    ) -> Tuple[bool, str, Optional[Path]]:
        """
        Safely update Master.xlsx using atomic replacement.
        Creates a backup before replacing if one exists.
        Validates the generated workbook before committing.
        """
        if df is None or not isinstance(df, pd.DataFrame):
            return False, "Validation failed: Master dataset must be a valid DataFrame.", None

        ensure_directories()
        backup_path = None

        if self.exists() and create_backup:
            try:
                backup_path = self.create_backup()
            except Exception as be:
                return False, f"Failed to create safety backup prior to update: {str(be)}", None

        # Temporary file in the same directory for true atomic replacement
        temp_file = self.master_path.parent / f"temp_{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}.xlsx"

        try:
            from .excel_exporter import write_master_excel_workbook
            write_master_excel_workbook(
                filepath=temp_file,
                cleaned_master_df=df,
                extra_sheets=extra_sheets,
            )

            # Verification step: Open temporary workbook and verify structure and content
            test_wb = openpyxl.load_workbook(temp_file, read_only=True)
            if "Cleaned_Master" not in test_wb.sheetnames:
                test_wb.close()
                temp_file.unlink(missing_ok=True)
                return False, "Validation error: 'Cleaned_Master' worksheet not found in generated workbook.", backup_path

            test_ws = test_wb["Cleaned_Master"]
            test_row_count = test_ws.max_row
            test_wb.close()

            # Ensure expected rows exist (header + data rows)
            expected_min_rows = 1 if df.empty else len(df) + 1
            if test_row_count is not None and test_row_count < expected_min_rows:
                temp_file.unlink(missing_ok=True)
                return False, f"Validation error: Generated master has {test_row_count} rows, expected at least {expected_min_rows}.", backup_path

            # Atomic replace (same directory)
            os.replace(str(temp_file), str(self.master_path))
            return True, "Master workbook saved successfully.", backup_path
        except Exception as e:
            if temp_file.exists():
                temp_file.unlink(missing_ok=True)
            # If master was somehow damaged and backup exists, recover
            if not self.exists() and backup_path and backup_path.exists():
                shutil.copy2(backup_path, self.master_path)
            return False, f"Failed to save Master workbook: {str(e)}", backup_path

    def restore_backup(self, backup_filename: str) -> Tuple[bool, str]:
        """Restore Master.xlsx from a specific backup file."""
        backup_path = self.backup_dir / backup_filename
        if not backup_path.exists():
            return False, f"Backup file '{backup_filename}' does not exist."

        # Create a pre-restore backup if master currently exists
        if self.exists():
            self.create_backup(prefix="Master_prerestore")

        try:
            shutil.copy2(backup_path, self.master_path)
            return True, f"Master workbook restored from '{backup_filename}'."
        except Exception as e:
            return False, f"Failed to restore backup: {str(e)}"

    def list_backups(self) -> List[Dict[str, Any]]:
        """List all available backup files ordered by most recent first."""
        if not self.backup_dir.exists():
            return []

        backups = []
        files = sorted(self.backup_dir.glob("Master_*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)
        for f in files:
            stat = f.stat()
            backups.append({
                "filename": f.name,
                "path": str(f),
                "size_kb": round(stat.st_size / 1024, 2),
                "created_at": datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M:%S"),
            })
        return backups

    def reset_master(self) -> Tuple[bool, str]:
        """Reset the master workbook after creating a safety backup."""
        if not self.exists():
            return True, "No master workbook existed to reset."

        backup = self.create_backup()
        try:
            self.master_path.unlink()
            return True, f"Master workbook reset. Previous state backed up to {backup.name if backup else 'N/A'}."
        except Exception as e:
            return False, f"Failed to reset master workbook: {str(e)}"

    def upload_custom_master(self, file_content_or_path: Any) -> Tuple[bool, str]:
        """Import an existing Master.xlsx provided by the user."""
        ensure_directories()
        if self.exists():
            self.create_backup()

        try:
            if hasattr(file_content_or_path, "read"):
                file_content_or_path.seek(0)
                with open(self.master_path, "wb") as f:
                    shutil.copyfileobj(file_content_or_path, f)
            else:
                shutil.copy2(str(file_content_or_path), self.master_path)

            # Validate that it can be opened
            df = self.load_master()
            if df is None:
                return False, "Uploaded file could not be read as a valid master workbook."
            return True, f"Master workbook successfully imported ({len(df)} rows)."
        except Exception as e:
            return False, f"Error importing master workbook: {str(e)}"
